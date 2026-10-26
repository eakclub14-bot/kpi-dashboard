import calendar
import io
import os
import re
import threading
import time
from collections import Counter
from datetime import datetime, timedelta
from flask import Flask, jsonify, render_template, request, send_file
import pandas as pd

try:
    import pythoncom
    import win32com.client
except ImportError:
    win32com = None
    pythoncom = None

try:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    pass

app = Flask(__name__)

LOADING_PROGRESS = 0

TARGET_ASSIGNEES = [
    "sasiprapa.ke@rsm.co.th",
    "panyakron.ch@rsm.co.th",
    "thapakorn.od@rsm.co.th",
    "husadin.sr@rsm.co.th",
    "nattapon.ni@rsm.co.th",
]

MONTH_NAME_TO_NUM = {
    'JAN': 1, 'JANUARY': 1, 'ม.ค.': 1, 'มกราคม': 1,
    'FEB': 2, 'FEBRUARY': 2, 'ก.พ.': 2, 'กุมภาพันธ์': 2,
    'MAR': 3, 'MARCH': 3, 'มี.ค.': 3, 'มีนาคม': 3,
    'APR': 4, 'APRIL': 4, 'เม.ย.': 4, 'เมษายน': 4,
    'MAY': 5, 'พ.ค.': 5, 'พฤษภาคม': 5,
    'JUN': 6, 'JUNE': 6, 'มิ.ย.': 6, 'มิถุนายน': 6,
    'JUL': 7, 'JULY': 7, 'ก.ค.': 7, 'กรกฎาคม': 7,
    'AUG': 8, 'AUGUST': 8, 'ส.ค.': 8, 'สิงหาคม': 8,
    'SEP': 9, 'SEPTEMBER': 9, 'ก.ย.': 9, 'กันยายน': 9,
    'OCT': 10, 'OCTOBER': 10, 'ต.ค.': 10, 'ตุลาคม': 10,
    'NOV': 11, 'NOVEMBER': 11, 'พ.ย.': 11, 'พฤศจิกายน': 11,
    'DEC': 12, 'DECEMBER': 12, 'ธ.ค.': 12, 'ธันวาคม': 12,
}

NUM_TO_MONTH_CODE = {
    1: 'JAN', 2: 'FEB', 3: 'MAR', 4: 'APR', 5: 'MAY', 6: 'JUN',
    7: 'JUL', 8: 'AUG', 9: 'SEP', 10: 'OCT', 11: 'NOV', 12: 'DEC',
}


DUE_DAY_OF_NEXT_MONTH = 5          # กำหนดส่ง: วันที่ 5 ของเดือนถัดไป
DUE_TIME = (23, 59, 59)         # นับถึงสิ้นวัน (เปลี่ยนเป็น (17, 0, 0) ถ้าต้องการตัดที่ 17:00)


def calculate_due_date_info(month_num, year, report_type):
    if not year or year < 2026:
        year = 2026

    # กำหนดให้ส่งวันที่ 5 ของเดือนถัดไปในปีเดิม (ยกเว้นเดือน 12 ที่จะไปเดือน 1 ปีถัดไป)
    if month_num == 12:
        next_month = 1
        next_year = year + 1
    else:
        next_month = month_num + 1
        next_year = year  # ปีเดิม ไม่ข้ามปี

    due_dt = datetime(next_year, next_month, DUE_DAY_OF_NEXT_MONTH, *DUE_TIME)

    due_str_iso = due_dt.strftime('%Y-%m-%d %H:%M:%S')
    due_str_excel = due_dt.strftime('%d/%m/%Y %H:%M:%S')

    return due_str_iso, due_str_excel, due_dt


def detect_report_details(subject, body, sent_date_str=None, monthend_hint=False):
    full_text = f"{subject} {body}".upper()          # ใช้หาเดือนของรายงาน (หัวเรื่องมาก่อนเนื้อหา)
    subject_text = (subject or "").upper()           # ใช้จัดหมวด: ดูจากหัวเรื่องเท่านั้น
    report_type = None
    category = "อื่นๆ"

    if re.search(r'CONTROL[\s_]*NUMBER', subject_text):
        report_type = 'CONTROL_NUMBER'
        category = 'Control Number'
    elif re.search(r'DSR', subject_text) and re.search(r'TAS[\s_]*HFS', subject_text):
        report_type = 'TAS_DSR'
        category = 'RSMGroup TAS DSR'
    elif re.search(r'DSR', subject_text) and (re.search(r'MONTH[\s_\-]*END', subject_text) or monthend_hint):
        # monthend_hint = หัวเรื่องไม่มีคำว่า Monthend แต่ชื่อไฟล์แนบมี (เช่น 'KKC_DSR MAS&HFS REPORT 31may 2026')
        report_type = 'MONTHEND_DSR'
        category = 'Monthend DSR'
    elif re.search(r'สินค้าแถม|ของแถม', subject_text):
        report_type = 'FREE_PREMIUM'
        category = 'รายงานสินค้าแถม'
    elif re.search(r'ยอดขาย|SALES', subject_text):
        report_type = 'SALES_REPORT'
        category = 'รายงานยอดขาย'

    detected_month = None
    detected_year = None

    month_regex = r'(JAN|JANUARY|ม\.ค\.|มกราคม|FEB|FEBRUARY|ก\.พ\.|กุมภาพันธ์|MAR|MARCH|มี\.ค\.|มีนาคม|APR|APRIL|เม\.ย\.|เมษายน|MAY|พ\.ค\.|พฤษภาคม|JUN|JUNE|มิ\.ย\.|มิถุนายน|JUL|JULY|ก\.ค\.|กรกฎาคม|AUG|AUGUST|ส\.ค\.|สิงหาคม|SEP|SEPTEMBER|ก\.ย\.|กันยายน|OCT|OCTOBER|ต\.ค\.|ตุลาคม|NOV|NOVEMBER|พ\.ย\.|พฤศจิกายน|DEC|DECEMBER|ธ\.ค\.|ธันวาคม)'
    year_regex = r"(?:\'|\s*\/|\s*)?(\d{4}|\d{2})?"

    matches = list(re.finditer(rf"{month_regex}{year_regex}", full_text, re.IGNORECASE))

    for match in matches:
        month_str = match.group(1).strip()
        year_str = match.group(2)

        temp_month = MONTH_NAME_TO_NUM.get(month_str)
        if temp_month:
            detected_month = temp_month
            if year_str:
                parsed_year = int(year_str)
                if parsed_year > 2500:
                    parsed_year -= 543
                elif parsed_year < 100:
                    parsed_year += 2000
                if parsed_year < 2026:
                    parsed_year = 2026
                detected_year = parsed_year
            break

    sent_dt = parse_date_string(sent_date_str) if sent_date_str else None

    if not detected_month and sent_dt:
        detected_month = sent_dt.month - 1 if sent_dt.month > 1 else 12
        detected_year = sent_dt.year if sent_dt.month > 1 else sent_dt.year - 1
    elif detected_month and not detected_year:
        detected_year = sent_dt.year if sent_dt else 2026

    if not detected_month:
        detected_month = 9

    if not detected_year or detected_year < 2026:
        detected_year = 2026

    period_tag = NUM_TO_MONTH_CODE.get(detected_month, 'SEP')
    due_iso, due_excel, due_dt = calculate_due_date_info(detected_month, detected_year, report_type)

    return due_iso, due_excel, due_dt, period_tag, report_type, category, detected_month, detected_year


def parse_date_string(date_str):
    if not date_str:
        return None
    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%d/%m/%Y %H:%M:%S',
        '%d/%m/%Y %H:%M',
    ]
    for fmt in formats:
        try:
            return datetime.strptime(str(date_str).strip(), fmt)
        except ValueError:
            pass
    return None


# ----------------------------- ค่าคงที่ --------------------------------------

FETCH_FROM_DATE = datetime(2026, 1, 1)   # ดึงเมลที่ได้รับตั้งแต่วันนี้เป็นต้นไป
# ช่วง "เดือนของรายงาน" ที่นับคะแนน (ม.ค.–ก.ย. 2026) ใช้กรอบเดียวกันทุกหมวด
MONTHEND_PERIOD_START = (2026, 1)
MONTHEND_PERIOD_END = (2026, 9)
BODY_LIMIT = 5000                        # อ่านเนื้อหาเมลไม่เกินกี่ตัวอักษร (เร็วขึ้นมาก)
OL_MAIL_ITEM = 43                        # olMail: ข้าม MeetingItem / Report / NDR

# ----------------------------- เกณฑ์คะแนน (ระบบจัดอันดับแบบ % ความสำเร็จ) -----------------
MAX_POINTS = 15          # คะแนนสูงสุดต่อรายงาน (ส่งเร็ว)
EARLY_HOURS = 48         # ส่งก่อนกำหนดอย่างน้อย 48 ชม. = ส่งเร็วเป็นพิเศษ (15)
POINTS_EARLY = 15
POINTS_ONTIME = 10       # ส่งทันกำหนดแต่ไม่ถึง 48 ชม.ก่อนกำหนด
POINTS_LATE = 5          # ส่งเลยกำหนด
# รายงานที่ "ควรส่ง" แต่ไม่ได้ส่ง = 0 คะแนน (นับรวมในตัวหาร จึงถูกลงโทษอัตโนมัติ)

# หมวดที่แต่ละคนรับผิดชอบ "ต่อเดือน" (ชื่อหมวดต้องตรงกับ category ใน detect_report_details)
EXPECTED_CATEGORIES = {
    "sasiprapa.ke@rsm.co.th": ["Monthend DSR", "Control Number"],
    "panyakron.ch@rsm.co.th": ["Monthend DSR", "Control Number"],
    "thapakorn.od@rsm.co.th": ["Monthend DSR", "Control Number"],
    "husadin.sr@rsm.co.th":   ["Monthend DSR", "Control Number", "รายงานสินค้าแถม"],
    "nattapon.ni@rsm.co.th":  ["Monthend DSR", "Control Number", "RSMGroup TAS DSR"],
}


def expected_periods(now=None):
    """เดือนของรายงาน (ม.ค.–ก.ย. 2026) ที่ครบกำหนดส่งไปแล้ว (ยังไม่ครบกำหนดจะไม่นับว่าขาดส่ง)"""
    now = now or datetime.now()
    periods = []
    y, m = MONTHEND_PERIOD_START
    while (y, m) <= MONTHEND_PERIOD_END:
        due = calculate_due_date_info(m, y, None)[2]
        if due < now:
            periods.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return periods


def score_submission(sent_dt, due_dt):
    """คืน (status_type, status_label, points)"""
    if sent_dt > due_dt:
        return "late", "ส่งล่าช้า", POINTS_LATE
    if (due_dt - sent_dt).total_seconds() >= EARLY_HOURS * 3600:
        return "early", "ส่งเร็วเป็นพิเศษ", POINTS_EARLY
    return "ontime", "ตรงเวลา", POINTS_ONTIME


# ถ้าชื่อที่แสดงใน Outlook ไม่ตรงกับอีเมล (เช่น ชื่อภาษาไทย) ให้ใส่ชื่อเล่น/ชื่อไทยตรงนี้
# ตัวอย่าง: "sasiprapa.ke@rsm.co.th": ["ศศิประภา"]
SENDER_NAME_ALIASES = {
    "sasiprapa.ke@rsm.co.th": [],
    "panyakron.ch@rsm.co.th": [],
    "thapakorn.od@rsm.co.th": [],
    "husadin.sr@rsm.co.th": [],
    "nattapon.ni@rsm.co.th": [],
}

# เลือกเมลจาก "หัวเรื่อง" เท่านั้น (ไม่ค้นในเนื้อหา) และต้องตรงรูปแบบรายงานที่กำหนด:
#   1) DSR Monthend: หัวเรื่องมี DSR และมีคำว่า Monthend / Month End
#   2) หัวข้ออื่นๆ:
#      - Control Number
#      - สินค้าแถม
#      - RSMGroup TAS HFS DSR (ต้องเจอคำว่า Monthend ด้วย: ในหัวเรื่อง หรือถ้าไม่มีให้หาในเนื้อหา/ชื่อไฟล์แนบ)
DSR_RE = re.compile(r'DSR', re.IGNORECASE)
MONTHEND_RE = re.compile(r'MONTH[\s_\-]*END', re.IGNORECASE)
TAS_RE = re.compile(r'TAS[\s_]*HFS', re.IGNORECASE)
CONTROL_RE = re.compile(r'CONTROL[\s_]*NUMBER', re.IGNORECASE)
FREE_RE = re.compile(r'สินค้าแถม|ของแถม')


def is_target_subject(subject):
    """หัวเรื่องผ่านเงื่อนไขทันที (ไม่ต้องดูเนื้อหา)"""
    s = subject or ""
    if DSR_RE.search(s) and MONTHEND_RE.search(s):      # DSR Monthend (รวม TAS ที่มี Monthend ในหัวเรื่อง)
        return True
    return bool(CONTROL_RE.search(s) or FREE_RE.search(s))


def needs_monthend_in_content(subject):
    """DSR ที่หัวเรื่องไม่มีคำว่า Monthend — ต้องไปหาคำว่า Monthend ต่อ
    - TAS HFS DSR: หาในเนื้อหา + ชื่อไฟล์แนบ
    - DSR อื่น (เช่น 'KKC_DSR MAS&HFS REPORT 31may 2026'): หาจาก "ชื่อไฟล์แนบ" เท่านั้น"""
    s = subject or ""
    return bool(DSR_RE.search(s) and not MONTHEND_RE.search(s))


def _attachment_names(msg):
    names = []
    try:
        for i in range(1, msg.Attachments.Count + 1):
            names.append(str(msg.Attachments.Item(i).FileName or ''))
    except Exception:
        pass
    return " ".join(names)


TARGET_LOWER = {e.lower().strip() for e in TARGET_ASSIGNEES}

# สถิติการดึงล่าสุด (ดูได้ที่ /api/fetch-stats)
LAST_FETCH_STATS = {}


# ----------------------------- ผู้ส่ง ----------------------------------------

def _match_assignee_by_name(sender_name):
    """เทียบชื่อที่แสดงกับทีม: ใช้ชื่อต้นของอีเมล (ก่อนจุด) เช่น sasiprapa.ke -> 'sasiprapa'"""
    name = (sender_name or "").lower().strip()
    if not name:
        return ""
    for assignee in TARGET_ASSIGNEES:
        key = assignee.lower()
        first = key.split('@')[0].split('.')[0]
        aliases = [a.lower() for a in SENDER_NAME_ALIASES.get(key, [])]
        if first in name or any(a in name for a in aliases):
            return key
    return ""


def get_clean_sender_email(msg):
    """คืนอีเมล SMTP ของผู้ส่ง หรือ "" ถ้าหาไม่ได้ (ไม่คืนค่า X.500 /O=EXCHANGELABS... แล้ว)"""
    # 1) SenderEmailAddress ที่เป็น SMTP อยู่แล้ว
    try:
        addr = str(msg.SenderEmailAddress or '').strip()
        if "@" in addr:
            return addr.lower()
    except Exception:
        pass

    # 2) Exchange user -> PrimarySmtpAddress
    try:
        sender_obj = msg.Sender
        if sender_obj:
            ex_user = sender_obj.GetExchangeUser()
            if ex_user and ex_user.PrimarySmtpAddress:
                return str(ex_user.PrimarySmtpAddress).lower().strip()
    except Exception:
        pass

    # 3) MAPI property ของ "ผู้ส่ง" บนตัวเมล (PR_SENDER_SMTP_ADDRESS)
    for tag in (
        "http://schemas.microsoft.com/mapi/proptag/0x5D01001F",  # PR_SENDER_SMTP_ADDRESS
        "http://schemas.microsoft.com/mapi/proptag/0x39FE001F",  # PR_SMTP_ADDRESS
    ):
        try:
            smtp = msg.PropertyAccessor.GetProperty(tag)
            if smtp and "@" in str(smtp):
                return str(smtp).lower().strip()
        except Exception:
            pass

    # 4) สุดท้ายเทียบจากชื่อที่แสดง
    try:
        return _match_assignee_by_name(str(msg.SenderName or ''))
    except Exception:
        return ""


# ----------------------------- ดึงเมล ----------------------------------------

# ผู้ส่งที่ให้ไปนับจาก "กล่องเมล/โฟลเดอร์อื่น" แทน Inbox หลัก   {อีเมลผู้ส่ง: ชื่อกล่องเมล/โฟลเดอร์}
# ถ้าหากล่องเมลนั้นไม่เจอ จะกลับไปนับจาก Inbox หลักตามเดิม (กันข้อมูลหาย) และแจ้งใน [fetch] alt_mailbox=...
ASSIGNEE_MAILBOX = {
    "husadin.sr@rsm.co.th": "เอกIT",
}
# กล่องเมลทางเลือกที่ต้องสแกนทั้งกล่อง จะข้ามโฟลเดอร์ที่ชื่อมีคำเหล่านี้
ALT_SKIP_WORDS = ("deleted", "junk", "sent", "draft", "outbox", "trash", "spam", "ลบ", "ขยะ", "ส่งแล้ว", "ร่าง")


def _norm_name(s):
    return re.sub(r'[\s_\-\.]+', '', str(s or '')).lower()


def _entry_id(folder):
    try:
        return folder.EntryID
    except Exception:
        return id(folder)


def _folder_name(folder):
    try:
        return str(folder.Name)
    except Exception:
        return "?"


def _iter_folders(folder, skip_words=()):
    """วนโฟลเดอร์นี้ + โฟลเดอร์ย่อยทั้งหมด (เผื่อมี Rule ย้ายเมลไปโฟลเดอร์ย่อย)
    skip_words: ข้ามโฟลเดอร์ย่อยที่ชื่อมีคำเหล่านี้ (ใช้กับกล่องเมลทางเลือก เช่น Sent/Deleted)"""
    yield folder
    try:
        for sub in folder.Folders:
            if skip_words:
                nm = _folder_name(sub).lower()
                if any(w in nm for w in skip_words):
                    continue
            yield from _iter_folders(sub, skip_words)
    except Exception:
        pass


def _find_mailbox_roots(namespace, wanted, max_depth=3):
    """หากล่องเมล (store) หรือโฟลเดอร์ที่ชื่อตรงกับ wanted (ไม่สนช่องว่าง/ตัวพิมพ์) คืนรายการโฟลเดอร์ที่ต้องสแกน"""
    target = _norm_name(wanted)
    roots, seen = [], set()

    def add(folder):
        eid = _entry_id(folder)
        if eid not in seen:
            seen.add(eid)
            roots.append(folder)

    def walk(folder, depth):
        try:
            subs = list(folder.Folders)
        except Exception:
            return
        for sub in subs:
            if target in _norm_name(_folder_name(sub)):
                add(sub)                                   # โฟลเดอร์ที่ชื่อตรง
            elif depth < max_depth:
                walk(sub, depth + 1)

    try:
        tops = list(namespace.Folders)
    except Exception:
        tops = []
    for top in tops:
        try:
            if target in _norm_name(_folder_name(top)):    # ทั้งกล่องเมลชื่อตรง
                try:
                    add(top.Store.GetDefaultFolder(6))     # ใช้ Inbox ของกล่องนั้น
                except Exception:
                    add(top)                               # เปิด Inbox ไม่ได้ -> สแกนทั้งกล่อง (ข้าม Sent/Deleted ฯลฯ)
            else:
                walk(top, 1)
        except Exception as e:
            print(f"[fetch] ค้นหากล่องเมล '{wanted}' ไม่สำเร็จในบางส่วน: {e}")
    return roots


def _process_message(msg, sender_cache, unmatched, only=None, exclude=frozenset()):
    """คืน (record, reason) — record เป็น dict ถ้าเก็บ, reason เป็นเหตุผลถ้าข้าม
    only: (กล่องเมลทางเลือก) นับเฉพาะผู้ส่งกลุ่มนี้   exclude: (Inbox หลัก) ข้ามผู้ส่งที่ไปนับกล่องเมลอื่นแล้ว"""
    try:
        if getattr(msg, "Class", None) != OL_MAIL_ITEM:
            return None, "not_mail"

        subject = str(msg.Subject or '')

        # เช็คหัวเรื่องก่อน (เรียก COM น้อยที่สุด) — ไม่ตรงรูปแบบก็ข้ามเลย ไม่โหลดผู้ส่ง/เนื้อหา
        content_check = None   # None / "tas" / "dsr"
        if not is_target_subject(subject):
            if needs_monthend_in_content(subject):
                content_check = "tas" if TAS_RE.search(subject) else "dsr"
            else:
                return None, ("dsr_no_monthend" if DSR_RE.search(subject) else "subject_mismatch")

        name = str(msg.SenderName or '')
        raw_addr = str(msg.SenderEmailAddress or '')
        cache_key = (name, raw_addr)
        if cache_key not in sender_cache:
            sender_cache[cache_key] = get_clean_sender_email(msg)
        sender = sender_cache[cache_key]

        if sender not in TARGET_LOWER:
            if only is None:
                unmatched[f"{name} | {raw_addr}"] += 1
                return None, "wrong_sender"
            return None, "alt_other_sender"
        if only is not None and sender not in only:
            return None, "alt_other_sender"
        if sender in exclude:
            return None, "moved_to_alt_mailbox"

        body = str(msg.Body or '')[:BODY_LIMIT]  # ใช้ช่วยหาเดือนของรายงานเท่านั้น ไม่ใช้คัดเลือกเมล

        # DSR ที่หัวเรื่องไม่มี Monthend: TAS -> เนื้อหา/ชื่อไฟล์แนบ, DSR อื่น -> ชื่อไฟล์แนบ
        attach_names = _attachment_names(msg)
        monthend_hint = False
        if content_check:
            haystack = f"{body} {attach_names}" if content_check == "tas" else attach_names
            if not MONTHEND_RE.search(haystack):
                return None, ("tas_no_monthend" if content_check == "tas" else "dsr_no_monthend")
            monthend_hint = True

        sent_date = msg.ReceivedTime.strftime('%Y-%m-%d %H:%M:%S')
        attachment_count = msg.Attachments.Count
        attachment_status = (
            f"มีไฟล์แนบ ({attachment_count} ไฟล์)"
            if attachment_count > 0
            else "ไม่มีไฟล์แนบ"
        )
        return {
            "subject": subject,
            "body": body,
            "assignee": sender,
            "sent_date": sent_date,
            "attachment_status": attachment_status,
            "attachment_names": attach_names,
            "monthend_hint": monthend_hint,
        }, None
    except Exception as e:
        print(f"[fetch] อ่านเมลไม่ได้: {e}")
        return None, "error"


# ---- ตัวกรองหัวเรื่องฝั่ง Outlook (DASL) : ให้ Outlook คัดมาก่อน ไม่ต้องวนอ่านทีละฉบับผ่าน COM ----
# ต้องกว้างกว่าหรือเท่ากับ is_target_subject() เสมอ (Python จะตรวจซ้ำอีกชั้น)
_SUBJ = '"urn:schemas:httpmail:subject"'
SUBJECT_DASL_FILTER = (
    "@SQL=("
    f"({_SUBJ} LIKE '%dsr%')"
    f" OR {_SUBJ} LIKE '%control%number%'"
    f" OR {_SUBJ} LIKE '%สินค้าแถม%'"
    f" OR {_SUBJ} LIKE '%ของแถม%'"
    ")"
)

# ---- Sync แบบเพิ่มเฉพาะเมลใหม่ : เก็บผลครั้งก่อนไว้ แล้วดึงเฉพาะเมลที่เข้ามาหลังจากนั้น ----
# เปิดได้เมื่อยืนยันแล้วว่าผลตรงกับการสแกนปกติ (เทียบ kept= ใน [fetch] ทั้งสองแบบ) — ปิดไว้เพื่อความถูกต้องก่อน
USE_DASL_PREFILTER = False

SYNC_LOCK = threading.Lock()                    # กัน sync ซ้อนกัน (เช่น กด Sync รัวๆ / Export ระหว่าง Sync)
EMAIL_CACHE = {"emails": [], "keys": set(), "last_sync": None}
SYNC_OVERLAP = timedelta(days=1)                # ย้อนเผื่อเมลเข้าช้า/นาฬิกาเหลื่อม


def _email_key(rec):
    return (rec["assignee"], rec["subject"], rec["sent_date"])


def _scan_outlook(since, full, use_dasl):
    """สแกน Outlook เฉพาะเมลที่ได้รับตั้งแต่ since
    คืน (emails, ok, suspicious) — suspicious=True เมื่อ filter ฝั่ง Outlook อาจใช้ไม่ได้ผล (ได้ 0 ทั้งที่มีเมลในช่วงวันที่)"""
    global LOADING_PROGRESS, LAST_FETCH_STATS
    t0 = time.time()
    LOADING_PROGRESS = 0
    emails = []
    stats = {
        "mode": "full" if full else "incremental",
        "since": since.strftime('%Y-%m-%d %H:%M'),
        "dasl_prefilter": use_dasl,
        "dasl_errors": 0,
        "alt_mailbox": {},
        "folders": 0,
        "folder_errors": 0,
        "checked": 0,
        "kept": 0,
        "duplicate": 0,
        "not_mail": 0,
        "wrong_sender": 0,
        "subject_mismatch": 0,
        "dsr_no_monthend": 0,
        "tas_no_monthend": 0,
        "moved_to_alt_mailbox": 0,
        "alt_other_sender": 0,
        "error": 0,
    }
    unmatched = Counter()
    ok = False
    suspicious = False

    if not win32com:
        LOADING_PROGRESS = 100
        LAST_FETCH_STATS = {**stats, "note": "ไม่พบ pywin32/Outlook"}
        return emails, False, False

    if pythoncom:
        pythoncom.CoInitialize()

    try:
        namespace = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
        inbox = namespace.GetDefaultFolder(6)

        # --- กล่องเมลทางเลือกของบางคน (เช่น husadin -> เอกIT) ---
        alt_roots = {}
        for assignee, mailbox_name in ASSIGNEE_MAILBOX.items():
            try:
                roots = _find_mailbox_roots(namespace, mailbox_name)
            except Exception as e:
                print(f"[fetch] หากล่องเมล '{mailbox_name}' ไม่ได้: {e}")
                roots = []
            stats["alt_mailbox"][assignee] = (
                f"{mailbox_name} -> " + ", ".join(_folder_name(r) for r in roots) if roots else f"{mailbox_name} NOT FOUND (ใช้ Inbox หลักแทน)"
            )
            if roots:
                alt_roots[assignee.lower()] = roots

        # รายการที่จะสแกน: (โฟลเดอร์, only, exclude)
        scan_list = []
        alt_ids = set()
        for assignee, roots in alt_roots.items():
            for root in roots:
                for f in _iter_folders(root, ALT_SKIP_WORDS):
                    eid = _entry_id(f)
                    if eid in alt_ids:
                        continue
                    alt_ids.add(eid)
                    scan_list.append((f, frozenset([assignee]), frozenset()))
        exclude_main = frozenset(alt_roots.keys())
        for f in _iter_folders(inbox):
            if _entry_id(f) in alt_ids:
                continue  # โฟลเดอร์นี้ถูกสแกนในโหมดกล่องเมลทางเลือกไปแล้ว
            scan_list.append((f, None, exclude_main))

        # กรองวันที่ที่ฝั่ง Outlook ก่อน (รูปแบบ US: MM/DD/YYYY hh:mm AM/PM)
        date_filter = f"[ReceivedTime] >= '{since.strftime('%m/%d/%Y %I:%M %p')}'"

        folder_items = []
        total = 0
        dated_total = 0
        for folder, only, exclude in scan_list:
            try:
                items = folder.Items
                items.Sort("[ReceivedTime]", True)
                dated = items.Restrict(date_filter)
                items = dated
                if use_dasl:
                    try:
                        items = dated.Restrict(SUBJECT_DASL_FILTER)
                        if full:
                            dated_total += dated.Count
                    except Exception as e:
                        stats["dasl_errors"] += 1
                        print(f"[fetch] DASL filter ใช้ไม่ได้ในโฟลเดอร์นี้ ใช้การสแกนปกติแทน: {e}")
                        items = dated
                count = items.Count
                folder_items.append((folder, items, only, exclude))
                total += count
                stats["folders"] += 1
            except Exception as e:
                stats["folder_errors"] += 1
                print(f"[fetch] เปิดโฟลเดอร์ไม่ได้: {e}")

        suspicious = bool(use_dasl and full and total == 0 and dated_total > 0)
        total = max(total, 1)
        sender_cache = {}
        seen = set()

        if not suspicious:
            for folder, items, only, exclude in folder_items:
                msg = items.GetFirst()
                while msg is not None:
                    stats["checked"] += 1
                    LOADING_PROGRESS = min(99, int(stats["checked"] / total * 100))

                    record, reason = _process_message(msg, sender_cache, unmatched, only, exclude)
                    if record:
                        key = _email_key(record)
                        if key in seen:
                            stats["duplicate"] += 1
                        else:
                            seen.add(key)
                            emails.append(record)
                            stats["kept"] += 1
                    elif reason:
                        stats[reason] += 1

                    try:
                        msg = items.GetNext()
                    except Exception:
                        break

        ok = stats["folders"] > 0

    except Exception as e:
        print(f"Error fetching Outlook emails: {e}")
        stats["fatal_error"] = str(e)
    finally:
        if pythoncom:
            pythoncom.CoUninitialize()
        LOADING_PROGRESS = 100

    stats["seconds"] = round(time.time() - t0, 2)
    stats["report_like_from_other_senders"] = unmatched.most_common(30)
    LAST_FETCH_STATS = stats
    print(
        "[fetch] "
        + ", ".join(f"{k}={v}" for k, v in stats.items() if k != "report_like_from_other_senders")
    )
    if unmatched:
        print("[fetch] เมลที่หัวเรื่องตรงรายงานแต่ผู้ส่งไม่ตรงทีม (ถ้ามีชื่อทีมปนอยู่ ให้เพิ่มใน SENDER_NAME_ALIASES):")
        for who, n in unmatched.most_common(30):
            print(f"        {n:4d}  {who}")

    return emails, ok, suspicious


def fetch_emails_from_outlook(full=False):
    """Sync เมลจาก Outlook
    - ครั้งแรก (หรือ full=True): สแกนทั้งช่วงตั้งแต่ FETCH_FROM_DATE
    - ครั้งต่อไป: ดึงเฉพาะเมลที่เข้ามาใหม่ แล้วรวมกับที่เก็บไว้ (เร็วกว่ามาก)
      หมายเหตุ: เมลที่ถูกลบ/ย้ายออกหลังจาก sync แล้วจะยังค้างอยู่จนกว่าจะสั่ง full (เรียก /api/dashboard?full=1)"""
    with SYNC_LOCK:
        started = datetime.now()
        incremental = (not full) and EMAIL_CACHE["last_sync"] is not None
        if incremental:
            since = max(FETCH_FROM_DATE, EMAIL_CACHE["last_sync"] - SYNC_OVERLAP)
        else:
            since = FETCH_FROM_DATE

        new_emails, ok, suspicious = _scan_outlook(since, not incremental, USE_DASL_PREFILTER)
        if suspicious:
            print("[fetch] DASL filter ไม่คืนเมลเลยทั้งที่มีเมลในช่วงวันที่ — สแกนปกติอีกรอบเพื่อความชัวร์")
            new_emails, ok, _ = _scan_outlook(since, not incremental, False)

        if ok:
            if not incremental:
                EMAIL_CACHE["emails"] = []
                EMAIL_CACHE["keys"] = set()
            for rec in new_emails:
                key = _email_key(rec)
                if key not in EMAIL_CACHE["keys"]:
                    EMAIL_CACHE["keys"].add(key)
                    EMAIL_CACHE["emails"].append(rec)
            EMAIL_CACHE["last_sync"] = started

        return list(EMAIL_CACHE["emails"])


def keep_latest_per_period(raw_emails):
    """รายงานเดียวกัน (ผู้ส่งคนเดียวกัน + หมวดเดียวกัน + เดือนของรายงานเดียวกัน) ที่ส่งมาหลายฉบับ
    เก็บเฉพาะ "ฉบับแรก" ที่ส่ง/รับ — ไม่ลงโทษคนที่ส่งทันแล้วส่งไฟล์แก้ไขทีหลังกำหนด
    (ชื่อฟังก์ชันคงเดิมเพื่อไม่ต้องแก้ที่เรียกใช้)"""
    first = {}
    for email in raw_emails:
        assignee = email.get("assignee", "Unknown").lower()
        sent_date_str = email.get("sent_date", "")
        (_, _, _, _, _, category, m_num, y_num) = detect_report_details(
            email.get("subject", ""), email.get("body", ""), sent_date_str,
            email.get("monthend_hint", False),
        )
        sent_dt = parse_date_string(sent_date_str) or datetime.max
        key = (assignee, category, y_num, m_num)
        if key not in first or sent_dt < first[key][0]:
            first[key] = (sent_dt, email)
    kept = [e for _, e in first.values()]
    return kept, len(raw_emails) - len(kept)


SAMPLE_EMAILS = [
    # --- SEP 2026 ---
    {"subject": "DSR_MAS&HFS REPORT 30 SEP'26(MonthEnd)", "body": "ส่งรายงาน UTH_MAS HFS DSR Tracking ประจำเดือน September 2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-10-01 10:48:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RSM MAS HFS DSR Tracking Form_30 Sep'26[MonthEnd]", "body": "ส่งรายงาน DSR ประจำเดือน September 2026 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-10-01 15:21:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS เดือน Sep 2026 RCB CD1: 11,439 CD2: 10,631 CD3: 9,914", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-10-05 09:50:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "RCB_MAS_HFS_DSR_Tracking [30 Sep 26] MonthEnd", "body": "ไฟล์แนบ CBR_MAS HFS DSR Tracking_V15 (30 Sep 26) - MonthEnd.xlsx", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-10-01 15:16:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "ข้อมูลสินค้าแถมประจำเดือน SEP'2026", "body": "ส่งรายงานของแถมประจำเดือน กันยายน 2569 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-10-01 15:11:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "BON_MAS_HFS_DSR_Tracking_30Sep26 (MonthEnd)", "body": "ไฟล์แนบ BON_MAS_HFS_DSR_Tracking_V15_30Sep26(MonthEnd).xlsx", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-10-01 14:04:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "Monthend KKC_DSR MAS&HFS REPORT 30 Sep 2026", "body": "ไฟล์แนบ Monthend KKC_MAS_HFS_DSR_Tracking_V15 30 Sep 26.xlsx", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-10-01 10:14:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "ข้อมูล Control Number และ Promotion IM ครับ", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-10-03 13:46:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},

    # --- AUG 2026 ---
    {"subject": "MONTHEND_DSR_MAS&HFS REPORT 31 AUG'26", "body": "ส่งรายงาน UTH_MAS_HFS_DSR_Tracking ประจำเดือน August 2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-09-01 11:11:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: RCB_MAS_HFS_DSR_Tracking [31 Aug 26] MonthEnd", "body": "CBR_MAS HFS DSR Tracking_V15 (31 Aug 26) - MonthEnd.xlsx", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-09-01 08:22:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "CD HFS เดือน Aug 2026 RCB CD1: 11,580 CD2: 11,578 CD3: 7,406", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-09-02 11:10:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "Monthend KKC_DSR MAS&HFS REPORT 31 Aug 2026", "body": "Monthend KKC_MAS_HFS_DSR_Tracking_V15 31 Aug 26.xlsx", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-09-01 10:58:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS) KKC", "body": "CD HFS เดือน Aug 2026 KKC CD1: 11,159 CD2: 12,157 CD3: 8,059", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-09-02 10:23:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "ข้อมูล Control Number และ Promotion IM Aug'26 ครับ", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-09-02 13:41:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "BON_MAS_HFS_DSR_Tracking_31Aug26 (MonthEnd)", "body": "ไฟล์แนบ BON_MAS_HFS_DSR_Tracking_V15_31Aug26(MonthEnd).xlsx", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-09-01 10:51:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: RSM MAS HFS DSR Tracking 31 Aug'26 [Month End]", "body": "ไฟล์แนบ RSM_MAS_HFS_DSR_Tracking_V15 Form (31 Aug'26)_MonthEND.xlsx", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-09-01 10:39:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS RSM AUG'26 ครับ CD1: 17,823 CD2: 17,661 CD3: 12,715", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-09-02 14:32:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "ส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS) UTH", "body": "ส่งคะแนน CD AUG'2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-09-02 10:49:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS) KKC", "body": "คะแนน CD AUG'2026 ครับ", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-09-02 10:22:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},

    # --- JUL 2026 ---
    {"subject": "Monthend KKC_DSR MAS&HFS REPORT 31 jul 2026", "body": "ส่งรายงาน Monthend KKC_MAS_HFS_DSR_Tracking ประจำเดือน July 2026 ครับ", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-08-01 10:49:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RSM MAS HFS DSR Tracking 31 Jul'26 [MonthEnd]", "body": "ส่งรายงาน RSM MAS HFS DSR Tracking ประจำเดือน July 2026 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-08-01 10:24:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RCB_MAS_HFS_DSR_Tracking [31 Jul 26] MonthEnd", "body": "ส่งรายงาน RCB_MAS_HFS_DSR_Tracking ประจำเดือน July 2026 ครับ", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-08-01 10:14:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS เดือน Jul 2026 RCB CD1: 11,094 CD2: 10,047 CD3: 7,525", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-08-03 13:17:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "DSR_MAS&HFS REPORT 31 JULY 26(MonthEnd)", "body": "ส่งรายงาน UTH_MAS_HFS_DSR_Tracking ประจำเดือน July 2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-08-01 08:23:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "BON_MAS_HFS_DSR_Tracking_31Jul26 (MonthEnd)", "body": "ไฟล์แนบ BON_MAS_HFS_DSR_Tracking_V15_31Jul26(MonthEnd).xlsx", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-08-01 11:35:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "เรียนพี่เอกพล ข้อมูล Control Number , CD และ Promotion IM ครับ", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-08-03 13:44:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS RSM JUL'26 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-08-03 13:24:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "ส่ง Control Number ตรวจสอบยอดขาย JUL 2026", "body": "คะแนน CD JUL'2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-08-04 10:15:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},

    # --- JUN 2026 ---
    {"subject": "Monthend KKC_DSR MAS&HFS REPORT 30 jun 2026", "body": "ส่งรายงาน Monthend KKC_MAS_HFS_DSR_Tracking ประจำเดือน June 2026 ครับ", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-07-01 10:45:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RSM MAS HFS DSR Tracking [30 Jun'26] - Month End", "body": "ส่งรายงาน RSM_MAS_HFS_DSR_Tracking ประจำเดือน June 2026 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-07-01 13:40:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RCB_MAS_HFS_DSR_Tracking [30 Jun 26] New DGP - MonthEnd", "body": "ส่งรายงาน RCB_MAS_HFS_DSR_Tracking ประจำเดือน June 2026 ครับ", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-07-01 12:01:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "DSR_MAS&HFS REPORT 30 June 26 GIV(MonthEnd)", "body": "ส่งรายงาน UTH_MAS_HFS_DSR_Tracking_V15 ประจำเดือน June 2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-07-01 10:50:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "เรียนพี่เอกพล ข้อมูล Control Number , CD และ Pro IM ครับ", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-07-03 08:23:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "BON_MAS_HFS_DSR_Tracking_30Jun26 (MonthEnd)", "body": "ไฟล์แนบ BON_MAS_HFS_DSR_Tracking_V15_30Jun26(MonthEnd).xlsx", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-07-01 10:19:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "ส่ง Control Number ตรวจสอบยอดขาย JUN 2026", "body": "คะแนน CD JUN'2026 ครับ", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-07-02 09:30:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},

    # --- APR, FEB, MAR, MAY, JAN 2026 ---
    {"subject": "RCB_MAS_HFS_DSR_Tracking [30 Apr 26] MonthEnd", "body": "ส่งรายงาน RCB_MAS_HFS_DSR_Tracking ประจำเดือน April 2026 ครับ", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-05-01 16:55:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS เดือน Mar 2026 KKC CD1: 11,432 CD2: 8,120 CD3: 9,116", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-04-04 08:18:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "monthend KKC_DSR MAS&HFS REPORT 31 Mar 2026", "body": "ส่งรายงาน Monthend KKC_MAS_HFS_DSR_Tracking ประจำเดือน March 2026 ครับ", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-04-01 11:03:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "KKC_DSR MAS&HFS REPORT 31may 2026", "body": "", "assignee": "thapakorn.od@rsm.co.th", "sent_date": "2026-06-02 08:57:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)", "attachment_names": "Monthend KKC_MAS_HFS_DSR_Tracking_V15 31 May 2026 New DGP.xlsx", "monthend_hint": True},
    {"subject": "RSM MAS HFS DSR Tracking 31 May'26 [MonthEnd]", "body": "ส่งรายงาน RSM_MAS_HFS_DSR_Tracking ประจำเดือน May 2026 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-06-01 09:10:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE: รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "คะแนน CD ร้านค้า HFS เดือน May 2026 RSM CD1: 17,472 CD2: 15,818 CD3: 9,588", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-06-02 13:56:00", "attachment_status": "ไม่มีไฟล์แนบ"},
    {"subject": "ส่งรายงานสินค้าแถมประจำเดือน MAY 2026 RSMGROUP", "body": "ข้อมูลของแถมประจำเดือน MAY'26 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-06-04 10:36:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "ส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS) UTH", "body": "ส่งคะแนน CD MAY'2026 ครับ", "assignee": "husadin.sr@rsm.co.th", "sent_date": "2026-06-04 08:44:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RE:รบกวนช่วยส่ง Control Number ตรวจสอบยอดขาย และ คะแนน CD (HFS)", "body": "ไฟล์ Promotion IM Feb26.xlsx", "assignee": "nattapon.ni@rsm.co.th", "sent_date": "2026-03-04 16:24:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "ส่ง Control Number ตรวจสอบยอดขาย JAN 2026", "body": "คะแนน CD JAN'2026 ครับ", "assignee": "panyakron.ch@rsm.co.th", "sent_date": "2026-02-04 09:15:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
    {"subject": "RCB_MAS_HFS_DSR_Tracking [30 Apr 26] MonthEnd", "body": "ส่งรายงาน RCB_MAS_HFS_DSR_Tracking ประจำเดือน April 2026 ครับ", "assignee": "sasiprapa.ke@rsm.co.th", "sent_date": "2026-05-01 16:55:00", "attachment_status": "มีไฟล์แนบ (1 ไฟล์)"},
]


def build_dashboard_payload(full=False):
    raw_emails = fetch_emails_from_outlook(full=full)

    use_sample = (
        os.environ.get("USE_SAMPLE_DATA") == "1"
        or (win32com is None and not raw_emails)
    )
    if use_sample:
        raw_emails = [dict(e) for e in SAMPLE_EMAILS]

    # รายงานซ้ำของเดือน/หมวดเดียวกัน -> เอาเฉพาะฉบับแรก
    raw_emails, older_dropped = keep_latest_per_period(raw_emails)
    LAST_FETCH_STATS["older_versions_dropped"] = older_dropped

    periods = set(expected_periods())   # เดือนของรายงานที่ครบกำหนดแล้ว

    processed_reports = []
    user_stats = {
        user.lower(): {"points": 0, "total": 0, "ontime": 0, "early": 0, "late": 0}
        for user in TARGET_ASSIGNEES
    }

    for email in raw_emails:
        subject = email.get("subject", "")
        body = email.get("body", "")
        assignee = email.get("assignee", "Unknown").lower()
        sent_date_str = email.get("sent_date", "")

        (
            due_iso,
            due_excel,
            due_dt,
            period_tag,
            report_type,
            category,
            m_num,
            y_num,
        ) = detect_report_details(
            subject, body, sent_date_str, email.get("monthend_hint", False)
        )

        # Monthend DSR นับเฉพาะเดือนของรายงานในช่วงที่กำหนด (ม.ค.–ก.ย. 2026)
        if report_type == 'MONTHEND_DSR' and not (
            MONTHEND_PERIOD_START <= (y_num, m_num) <= MONTHEND_PERIOD_END
        ):
            continue

        sent_dt = parse_date_string(sent_date_str)

        status_type = "ontime"
        status_label = "ตรงเวลา"
        points = POINTS_ONTIME
        delay_reason = "-"

        if sent_dt and due_dt:
            status_type, status_label, points = score_submission(sent_dt, due_dt)
            if status_type == "late":
                late_days = max(1, (sent_dt - due_dt).days)
                delay_reason = f"ส่งเกินกำหนด (ล่าช้าประมาณ {late_days} วัน)"
            elif status_type == "early":
                hours = int((due_dt - sent_dt).total_seconds() // 3600)
                delay_reason = f"ส่งก่อนกำหนด (เร็วกว่าประมาณ {hours} ชม.)"
            else:
                delay_reason = "ส่งภายใน 48 ชม. ก่อนกำหนด"

        # นับคะแนนเฉพาะ: ผู้ส่งในทีม + หมวดที่คนนั้นรับผิดชอบ + เดือนที่ครบกำหนดแล้ว
        counted = (
            assignee in user_stats
            and category in EXPECTED_CATEGORIES.get(assignee, [])
            and (y_num, m_num) in periods
        )
        if counted:
            s = user_stats[assignee]
            s["total"] += 1
            s["points"] += points
            if status_type in ("ontime", "early"):
                s["ontime"] += 1
            if status_type == "early":
                s["early"] += 1
            if status_type == "late":
                s["late"] += 1

        processed_reports.append({
            "report_name": subject,
            "subject": subject,
            "assignee": assignee,
            "submitted_date": sent_date_str,
            "sent_date": sent_date_str,
            "due_date": due_iso,
            "due_date_display": due_iso,
            "due_date_excel": due_excel,
            "period": period_tag,
            "category": category,
            "status": status_label,
            "status_type": status_type,
            "points": int(points),
            "counted": counted,
            "attachment_status": email.get("attachment_status", "ไม่มีไฟล์แนบ"),
            "delay_reason": delay_reason,
        })

    # ----------------------------- จัดอันดับ ------------------------------------
    # คะแนนจัดอันดับ = คะแนนที่ได้ / (จำนวนรายงานที่ "ควรส่ง" x คะแนนเต็ม) x 100
    # รายงานที่ควรส่งแต่ไม่ส่ง = 0 คะแนน (อยู่ในตัวหาร) จึงเทียบกันได้แม้ภาระงาน/หมวดไม่เท่ากัน
    def achievement(email, s):
        expected = len(EXPECTED_CATEGORIES.get(email, [])) * len(periods)
        pct = (s["points"] / (expected * MAX_POINTS) * 100) if expected else 0.0
        return pct, expected

    ranked = sorted(
        user_stats.items(),
        key=lambda x: (
            achievement(*x)[0],                                           # 1) % ความสำเร็จ
            (x[1]["ontime"] / x[1]["total"]) if x[1]["total"] else 0,     # 2) อัตราส่งทันกำหนด
            x[1]["total"],                                                # 3) จำนวนรายงานที่ส่ง
        ),
        reverse=True,
    )

    leaderboard = []
    for rank, (email, s) in enumerate(ranked, 1):
        score_pct, expected = achievement(email, s)
        avg_points = (s["points"] / s["total"]) if s["total"] else 0.0
        ontime_rate = (s["ontime"] / s["total"] * 100) if s["total"] else 0.0
        leaderboard.append({
            "rank": rank,
            "assignee": email,
            "points": int(s["points"]),
            "avg_points": f"{avg_points:.1f}",
            "fair_score": f"{score_pct:.1f}",          # = % ความสำเร็จ (0-100)
            "ontime_rate": f"{ontime_rate:.1f}%",
            "total": s["total"],
            "ontime": s["ontime"],
            "late": s["late"],
            "expected": expected,
            "missing": max(0, expected - s["total"]),
        })

    return {"reports": processed_reports, "leaderboard": leaderboard}


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/progress')
def get_progress():
    global LOADING_PROGRESS
    return jsonify({"progress": LOADING_PROGRESS})


@app.route('/api/fetch-stats')
def get_fetch_stats():
    return jsonify(LAST_FETCH_STATS)


@app.route('/api/dashboard', methods=['GET'])
def get_dashboard_data():
    # /api/dashboard?full=1 = สแกนใหม่ทั้งหมด (ปกติจะดึงเฉพาะเมลใหม่)
    payload = build_dashboard_payload(full=request.args.get('full') == '1')
    return jsonify(payload)


@app.route('/api/export-excel', methods=['GET'])
def export_excel():
    dashboard_data = build_dashboard_payload()
    reports_data = dashboard_data.get("reports", [])
    leaderboard_data = dashboard_data.get("leaderboard", [])

    df_reports = pd.DataFrame(reports_data)
    df_leaderboard = pd.DataFrame(leaderboard_data)

    if not df_reports.empty:
        df_reports['due_date'] = df_reports['due_date_excel']
        df_reports['counted'] = df_reports['counted'].map({True: 'นับคะแนน', False: 'ไม่นับ'})
        df_reports = df_reports[
            [
                'subject',
                'assignee',
                'period',
                'category',
                'sent_date',
                'due_date',
                'status',
                'points',
                'counted',
                'attachment_status',
            ]
        ]
        df_reports.columns = [
            'ชื่อรายงาน',
            'ผู้รับผิดชอบ',
            'ประจำเดือน',
            'หมวดหมู่',
            'วันที่ส่ง',
            'วันกำหนดส่ง',
            'สถานะ',
            'คะแนน',
            'การนับคะแนน',
            'ไฟล์แนบ',
        ]

    if not df_leaderboard.empty:
        df_leaderboard = df_leaderboard[
            [
                'rank', 'assignee', 'points', 'avg_points', 'ontime_rate',
                'total', 'ontime', 'late', 'expected', 'missing', 'fair_score',
            ]
        ]
        df_leaderboard.columns = [
            'อันดับ',
            'ผู้รับผิดชอบ',
            'คะแนนรวม',
            'คะแนนเฉลี่ย/รายงาน',
            'อัตราส่งตรงเวลา',
            'ส่งแล้ว (ฉบับ)',
            'ส่งทันกำหนด (ฉบับ)',
            'ส่งล่าช้า (ฉบับ)',
            'ควรส่ง (ฉบับ)',
            'ขาดส่ง (ฉบับ)',
            'คะแนนจัดอันดับ (% ความสำเร็จ)',
        ]

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_leaderboard.to_excel(
            writer, sheet_name='Leaderboard', index=False, startrow=3
        )
        df_reports.to_excel(
            writer, sheet_name='รายงานทั้งหมด', index=False, startrow=3
        )

        workbook = writer.book

        navy_header_fill = PatternFill(
            start_color="1F4E79", end_color="1F4E79", fill_type="solid"
        )
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        title_font = Font(name="Segoe UI", size=15, bold=True, color="1F4E79")
        data_font = Font(name="Segoe UI", size=10)

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9'),
        )

        status_ontime_fill = PatternFill(
            start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
        )
        status_ontime_font = Font(
            name="Segoe UI", size=10, color="375623", bold=True
        )

        status_late_fill = PatternFill(
            start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
        )
        status_late_font = Font(
            name="Segoe UI", size=10, color="C65911", bold=True
        )

        ws_lb = workbook['Leaderboard']
        ws_lb['A1'] = "📊 รายงานสรุปอันดับผู้ส่งรายงานประจำปี 2026 (Fair Ranking: % ความสำเร็จ)"
        ws_lb['A1'].font = title_font
        ws_lb['A2'] = (
            f"คะแนนจัดอันดับ = คะแนนที่ได้ ÷ (จำนวนรายงานที่ควรส่ง × {MAX_POINTS}) × 100  |  "
            "รายงานที่ควรส่งแต่ไม่ส่ง = 0 คะแนน"
        )
        ws_lb['A2'].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

        ws_rep = workbook['รายงานทั้งหมด']
        ws_rep['A1'] = "📋 รายละเอียดการส่งรายงานทั้งหมดประจำปี 2026"
        ws_rep['A1'].font = title_font

        for ws in [ws_lb, ws_rep]:
            ws.views.sheetView[0].showGridLines = True

            ws.row_dimensions[4].height = 28
            for cell in ws[4]:
                cell.fill = navy_header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
                ws.row_dimensions[row[0].row].height = 22
                for cell in row:
                    cell.border = thin_border
                    cell.font = data_font
                    cell.alignment = Alignment(
                        vertical="center", horizontal="left"
                    )

                    val_str = str(cell.value or '').strip()

                    if ws == ws_lb and cell.column_letter != 'B':
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center"
                        )
                    elif ws == ws_rep and cell.column_letter not in ['A', 'B']:
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center"
                        )

                    if ws == ws_rep:
                        if "ตรงเวลา" in val_str or "ส่งเร็ว" in val_str:
                            cell.fill = status_ontime_fill
                            cell.font = status_ontime_font
                            cell.alignment = Alignment(
                                horizontal="center", vertical="center"
                            )
                        elif val_str == "ส่งล่าช้า":
                            cell.fill = status_late_fill
                            cell.font = status_late_font
                            cell.alignment = Alignment(
                                horizontal="center", vertical="center"
                            )

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row >= 4 and cell.value:
                        val = str(cell.value)
                        length = sum(
                            2 if ord(char) > 255 else 1 for char in val
                        )
                        if length > max_len:
                            max_len = length
                adjusted_width = min(max(max_len + 4, 12), 60)
                ws.column_dimensions[col_letter].width = adjusted_width

    output.seek(0)

    return send_file(
        output,
        download_name="Report_Tracking_2026.xlsx",
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)